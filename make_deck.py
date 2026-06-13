import collections.abc  # noqa: F401
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt


def add_styled_slide(prs, title, content):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = DARK_BG

    title_box = slide.shapes.add_textbox(Inches(0.75), Inches(0.55), Inches(11.8), Inches(1.1))
    title_p = title_box.text_frame.paragraphs[0]
    title_p.text = title
    title_p.font.name = "Arial"
    title_p.font.size = Pt(34)
    title_p.font.bold = True
    title_p.font.color.rgb = WHITE

    body_box = slide.shapes.add_textbox(Inches(0.75), Inches(1.8), Inches(11.8), Inches(5.35))
    body_tf = body_box.text_frame
    body_tf.word_wrap = True
    body_tf.margin_left = Inches(0.05)
    body_tf.margin_right = Inches(0.05)

    for idx, line in enumerate(content.split("\n")):
        p = body_tf.paragraphs[0] if idx == 0 else body_tf.add_paragraph()
        p.text = line
        p.font.name = "Arial"

        emphasized = any(
            token in line
            for token in (
                "->",
                "$",
                "₹",
                "TAM",
                "SAM",
                "SOM",
                "Level 1",
                "Level 99",
                "RiseWise",
                "Gold Tier",
                "Seeking",
            )
        )
        if emphasized:
            p.font.size = Pt(23)
            p.font.bold = True
            p.font.color.rgb = NEON_GREEN
        else:
            p.font.size = Pt(17)
            p.font.color.rgb = LIGHT_GREY
        p.space_after = Pt(7)


prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)

DARK_BG = RGBColor(13, 17, 23)
WHITE = RGBColor(255, 255, 255)
NEON_GREEN = RGBColor(0, 230, 118)
LIGHT_GREY = RGBColor(180, 185, 195)

slides_list = [
    (
        "Slide 1: The Title Slide",
        "RiseWise\n"
        "Redefining the Piggy Bank for Gen Z and Gen Alpha in India\n\n"
        "Core Pillars:\n"
        "• The Habit: Micro-savings that compound into massive wealth.\n"
        "• The Tech: Gamified learning modules, live quizzes, and streak metrics.\n"
        "• The Edge: Seamless progression from basic savings to advanced options."
    ),
    (
        "Slide 2: The Philosophy - The Digital Gullak",
        "The Power of Everyday Micro-Savings\n"
        "The Concept: Modernizing the traditional clay piggy bank used by earlier generations.\n"
        "₹100/day -> ₹2 Crore to ₹4 Crore\n\n"
        "Compound Breakdown:\n"
        "• Skipping one premium tea or delivery can save ₹100 daily (₹3,000/month).\n"
        "• Invested from age 15 to late 40s across index funds and equity assets.\n"
        "• At 12% to 15% CAGR, this can compound into roughly ₹1.95 Crore to ₹4.46 Crore."
    ),
    (
        "Slide 3: The Engagement - Gamified Ecosystem",
        "Built Like a Video Game, Not a Textbook\n\n"
        "Daily Quizzes and Streaks:\n"
        "• Bite-sized 60-second finance challenges.\n"
        "• XP rewards and streak loops for daily retention.\n\n"
        "Leaderboards and Social Competition:\n"
        "• Global and private portfolio leagues.\n"
        "• Badges and profile upgrades as users level up.\n\n"
        "Risk-Free Simulations:\n"
        "• Virtual-currency simulations mapped to live Indian markets."
    ),
    (
        "Slide 4: Curriculum Evolution and Risk Profiles",
        "A Curriculum That Grows With the User\n\n"
        "Level 1: Foundational Free Tier\n"
        "• Micro-investing automation and inflation fundamentals.\n"
        "• Conservative risk profiles to preserve starter capital.\n\n"
        "Level 99: Advanced Trader Mastery\n"
        "• Options mechanics and professional risk-hedging concepts.\n"
        "• Defined-risk vertical spreads (Bull Put and Bear Call).\n"
        "• Delta, theta, and probability-of-profit tracking."
    ),
    (
        "Slide 5: Monetisation Pillar 1 - B2B Partner Ecosystem",
        "The Freemium Engine: Native Partner Placements\n\n"
        "Highly Targeted Audience:\n"
        "• Connects Mutual Fund houses and SEBI-registered brokers to high-intent youth.\n\n"
        "Subtle and Non-Intrusive:\n"
        "• No pop-ups, no disruptive banners, only contextual actions.\n\n"
        "Contextual Integration:\n"
        "• Quiz mastery -> native CTA to start an SIP with partner AMCs.\n"
        "• Advanced options mastery -> seamless redirect to partner broker demat onboarding."
    ),
    (
        "Slide 6: Monetisation Pillar 2 - The Gold Tier",
        "The Gold Tier: Institutional Tools for Retail Traders\n\n"
        "Value Proposition:\n"
        "• Paid layer for advanced users transitioning into active trading.\n\n"
        "Premium Feature Suite:\n"
        "• Real-time bulk and block deal visibility with FII/DII activity mapping.\n"
        "• Momentum alerts for large-volume spikes and accumulation footprints.\n"
        "• Backtesting for multi-leg options strategies with premium data feeds."
    ),
    (
        "Slide 7: Competitive Landscape - Market Players",
        "Feature Matrix (Condensed):\n"
        "• Primary Audience: RiseWise (Gen Z/Alpha) | Brokers (Working Professionals) | "
        "Kid Fintech (Ages 11-17) | Creators (Mass Audience)\n"
        "• Gamified XP Quizzes: RiseWise (Yes) | Others (No)\n"
        "• Risk-Free Derivatives Sandbox: RiseWise (Yes) | Others (No)\n"
        "• Institutional Bulk Data: RiseWise (Gold Tier) | Others (Limited/None)\n"
        "• Monetisation: RiseWise (Freemium + B2B Lead Gen) vs competitors' single-engine models"
    ),
    (
        "Slide 8: Market Sizing - The Indian Youth TAM",
        "Capturing India's Premium Youth Demographic\n\n"
        "TAM: 110 Million+ Users\n"
        "• Derived from affluent, connected segments within India's Gen Z and Gen Alpha populations.\n\n"
        "SAM: 45 Million Users\n"
        "• Smartphone-penetrated urban youth actively consuming fintech and gaming content.\n\n"
        "SOM: 3.5 Million Users\n"
        "• Year 3 goal for high-intent users transitioning from savings to active trading behaviors."
    ),
    (
        "Slide 9: The Investment Ask",
        "Seeking $5 Million for a 20% Equity Stake\n"
        "Implied Post-Money Valuation: $25 Million (~₹208 Crore)\n\n"
        "Capital Deployment:\n"
        "• 40% Product and Engineering: low-latency options data and risk engines.\n"
        "• 30% Gamification Pipeline: simulation leagues, adaptive quizzes, level systems.\n"
        "• 30% User Acquisition: targeted Tier-1 and Tier-2 campus growth campaigns."
    ),
    (
        "Slide 10: Founding Team",
        "Founder Spotlight\n\n"
        "Manoj Thakkar - CEO & Founder\n"
        "LinkedIn: https://www.linkedin.com/in/manoj-thakkar-6583b46/\n\n"
        "Leadership Focus:\n"
        "• Product vision for India's youth-first wealth platform.\n"
        "• Execution across gamification, fintech partnerships, and growth."
    ),
]

for slide_title, slide_content in slides_list:
    add_styled_slide(prs, slide_title, slide_content)

prs.save("RiseWise_Full_Investor_Deck.pptx")
print("Pitch deck presentation generated successfully!")
